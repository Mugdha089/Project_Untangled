"""Regression tests for live-upload sample isolation and evidence-based tags."""

from __future__ import annotations

import io
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import httpx
import networkx as nx
from openpyxl import Workbook
from openpyxl import load_workbook

from excel_parser import load_excel_tags
from graph_engine import (
    _point_on_segment,
    _build_equipment_nodes,
    apply_ambiguity_resolution,
    build_process_graph,
)
from main import app, lifespan, _source_tag_detections
import run_storage
import pdf_processor
from pdf_processor import (
    PdfProcessError,
    has_valid_source_bbox,
    match_excel_inventory,
    process_pdf,
    _extract_embedded_text,
)
from topology_evidence import CONNECTED_JUNCTION, find_ambiguous_process_crossings

def _find_test_cases_dir() -> Path:
    current = Path(__file__).resolve()
    for parent in (current.parent, *current.parents):
        candidate = parent / "test_cases"
        if candidate.exists():
            return candidate
    return current.parents[1] / "test_cases"


TEST_CASES = _find_test_cases_dir()
EXCEL_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
DEMO_EQUIPMENT_TAGS = {"TK-101", "P-101A", "P-101B", "E-102", "V-102"}


def _fixture_files(case_name: str) -> tuple[Path, Path]:
    pdf_path = TEST_CASES / f"{case_name}.pdf"
    excel_path = TEST_CASES / f"{case_name}_io.xlsx"
    return pdf_path, excel_path


def _supplied_fixture_files(case_name: str) -> tuple[Path, Path]:
    pack_dir = (
        Path(__file__).resolve().parents[2]
        / "UNTANGLE_mid_complex_test_pack"
        / "UNTANGLE_mid_complex_test_pack"
        / "test_cases"
    )
    return pack_dir / f"{case_name}.pdf", pack_dir / f"{case_name}_io.xlsx"


def _workbook_bytes(rows: list[tuple[str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Tag", "Type"])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class PdfEvidenceTests(unittest.TestCase):
    def test_inventory_sources_include_equipment_only_with_a_real_bbox(self) -> None:
        detections = _source_tag_detections(
            {
                "tags": [],
                "equipment": [
                    {"tag": "P-1", "type": "pump", "source_bbox": [10, 20, 30, 40]},
                    {"tag": "E-2", "type": "exchanger", "source_bbox": [0, 0, 0, 0]},
                ],
            }
        )
        self.assertEqual([item["tag_name"] for item in detections], ["P-1"])
        self.assertEqual(detections[0]["source_bbox"], [10, 20, 30, 40])

    def test_detected_instrument_is_not_assigned_to_nearest_equipment(self) -> None:
        fitz = pdf_processor.fitz
        if fitz is None:
            self.skipTest("PyMuPDF is unavailable")
        document = fitz.open()
        page = document.new_page()
        page.insert_text((48, 72), "PT-999")
        page.insert_text((470, 620), "P-999")
        pdf_bytes = document.tobytes()
        document.close()

        result = process_pdf(pdf_bytes, filename="unassociated-tag.pdf")

        tag = next(row for row in result["tags"] if row["tag_name"] == "PT-999")
        self.assertIsNone(tag["parent_equipment"])

    def test_graph_does_not_attach_missing_or_unassociated_instrument_tags(self) -> None:
        nodes = _build_equipment_nodes(
            [{"tag": "P-101", "type": "pump", "source_bbox": [10, 10, 20, 20]}],
            [
                {"tag_name": "PT-101", "status": "Matched", "parent_equipment": None},
                {"tag_name": "FT-101", "status": "Missing in P&ID", "parent_equipment": "P-101"},
            ],
        )
        self.assertEqual(nodes[0]["collapsed_tags"], [])


class BackendStartupConfigurationTests(unittest.IsolatedAsyncioTestCase):
    async def test_backend_dotenv_is_loaded_before_database_initialization(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            env_path = Path(directory) / "test-only.env"
            database_path = str(Path(directory) / "run-test.sqlite3")
            env_path.write_text(f"UNTANGLE_DB_PATH={database_path}\n", encoding="utf-8")
            with patch.dict(os.environ):
                os.environ.pop("UNTANGLE_DB_PATH", None)
                with patch("main.BACKEND_ENV_FILE", env_path), patch(
                    "main.run_storage.initialize_database",
                    side_effect=lambda: self.assertEqual(
                        os.environ.get("UNTANGLE_DB_PATH"), database_path
                    ),
                ) as initialize_database:
                    async with lifespan(app):
                        pass
                    initialize_database.assert_called_once_with()

    def test_unrecognized_equipment_tag_uses_explicit_nearby_label(self) -> None:
        class FakePage:
            def get_drawings(self):
                return []

            def get_text(self, _mode):
                return {
                    "blocks": [
                        {
                            "type": 0,
                            "lines": [
                                {"spans": [{"text": "MX-101", "bbox": (100, 100, 130, 110)}]},
                                {"spans": [{"text": "Mixer", "bbox": (102, 112, 125, 122)}]},
                                {"spans": [{"text": "PT-101", "bbox": (160, 100, 190, 110)}]},
                                {"spans": [{"text": "Pump", "bbox": (162, 112, 185, 122)}]},
                            ],
                        }
                    ]
                }

        class FakeDocument:
            def __len__(self):
                return 1

            def __getitem__(self, index):
                if index != 0:
                    raise IndexError(index)
                return FakePage()

        tags, equipment, _, _ = _extract_embedded_text(FakeDocument())

        self.assertIn(("MX-101", "mixer"), [(row["tag"], row["type"]) for row in equipment])
        self.assertIn("PT-101", [row["tag_name"] for row in tags])
        self.assertNotIn("PT-101", [row["tag"] for row in equipment])
        self.assertEqual(next(row for row in equipment if row["tag"] == "MX-101")["source_bbox"], [100, 100, 130, 110])

    def test_endpoint_tolerance_is_not_applied_as_a_line_length_fraction(self) -> None:
        self.assertFalse(_point_on_segment((150.0, 0.0), (0.0, 0.0), (100.0, 0.0), 2.0))
        self.assertTrue(_point_on_segment((100.75, 0.0), (0.0, 0.0), (100.0, 0.0), 1.0))

    def test_embedded_title_without_detectable_tags_is_not_reported_as_processed(self) -> None:
        fitz = pdf_processor.fitz
        if fitz is None:
            self.skipTest("PyMuPDF is unavailable")

        document = fitz.open()
        page = document.new_page()
        page.insert_text((48, 72), "P&ID scan sheet")
        pdf_bytes = document.tobytes()
        document.close()

        with self.assertRaisesRegex(PdfProcessError, "OCR is not configured"):
            process_pdf(pdf_bytes, filename="scanned-with-title.pdf")

    def test_matching_requires_valid_source_coordinates(self) -> None:
        inventory = [
            {"tag_name": "PT-101", "type": "Pressure", "excel_row": 2},
            {"tag_name": "PT-102", "type": "Pressure", "excel_row": 3},
            {"tag_name": "PT-103", "type": "Pressure", "excel_row": 4},
        ]
        detections = [
            {"tag_name": "PT-101", "source_bbox": [10, 20, 30, 40]},
            {"tag_name": "PT-102", "source_bbox": [0, 0, 0, 0]},
            {"tag_name": "PT-103", "source_bbox": [30, 40, 10, 20]},
        ]

        audit = match_excel_inventory(detections, inventory)

        self.assertEqual(
            [(row["tag_name"], row["status"]) for row in audit],
            [
                ("PT-101", "Matched"),
                ("PT-102", "Missing in P&ID"),
                ("PT-103", "Missing in P&ID"),
            ],
        )

    def test_demo_entities_are_available_only_when_explicitly_requested(self) -> None:
        pdf_path, _ = _fixture_files("02_parallel_pumps")
        inventory = load_excel_tags(
            _fixture_files("02_parallel_pumps")[1].read_bytes()
        )["tags"]

        live_result = process_pdf(pdf_path.read_bytes(), excel_tags=inventory)
        demo_result = process_pdf(pdf_path.read_bytes(), demo_mode=True)

        live_tags = {item["tag"] for item in live_result["equipment"]}
        self.assertTrue(DEMO_EQUIPMENT_TAGS.isdisjoint(live_tags), live_tags)
        self.assertIn("TK-101", {item["tag"] for item in demo_result["equipment"]})

    def test_fixture_equipment_detection_uses_pdf_tags_and_symbols(self) -> None:
        for case_name in (
            "01_easy_linear",
            "02_parallel_pumps",
            "03_branch_merge",
            "04_control_utility",
            "05_dense_recycle",
            "07_missing_ambiguous_tags",
        ):
            with self.subTest(case=case_name):
                pdf_path, excel_path = _fixture_files(case_name)
                result = process_pdf(pdf_path.read_bytes(), filename=pdf_path.name)
                workbook = load_workbook(
                    excel_path, data_only=True, read_only=True
                )
                sheet = workbook["Major Equipment"]
                rows = list(sheet.values)
                header = [str(value) for value in rows[0]]
                tag_column = header.index("Equipment Tag")
                expected_column = header.index("Expected in PDF")
                expected = {
                    str(row[tag_column]).strip()
                    for row in rows[1:]
                    if str(row[expected_column]).strip().lower() == "yes"
                }
                actual = {item["tag"] for item in result["equipment"]}
                self.assertEqual(actual, expected)
                workbook.close()

    def test_scanned_fixture_remains_explicitly_unsupported(self) -> None:
        pdf_path, _ = _fixture_files("06_scanned_image_only")
        with self.assertRaisesRegex(PdfProcessError, "OCR is not configured"):
            process_pdf(pdf_path.read_bytes(), filename=pdf_path.name)


class PipelineRegressionTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self.database_dir = tempfile.TemporaryDirectory()
        self.database_env = patch.dict(
            os.environ,
            {"UNTANGLE_DB_PATH": str(Path(self.database_dir.name) / "test.sqlite3")},
        )
        self.database_env.start()
        run_storage.initialize_database()
        self.client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app),
            base_url="http://test",
        )
        login = await self.client.post(
            "/api/auth/login",
            json={"email": "demo@untangle.local", "password": "untangle-demo-2026"},
        )
        self.assertEqual(login.status_code, 200, login.text)
        self.client.headers["Authorization"] = f"Bearer {login.json()['access_token']}"

    async def asyncTearDown(self) -> None:
        await self.client.aclose()
        self.database_env.stop()
        self.database_dir.cleanup()

    async def _process_case(self, case_name: str, *, demo_mode: bool = False) -> httpx.Response:
        pdf_path, excel_path = _fixture_files(case_name)
        return await self.client.post(
            "/api/process-pid",
            files={
                "file_pdf": (
                    pdf_path.name,
                    pdf_path.read_bytes(),
                    "application/pdf",
                ),
                "file_excel": (
                    excel_path.name,
                    excel_path.read_bytes(),
                    EXCEL_CONTENT_TYPE,
                ),
            },
            data={"demo_mode": str(demo_mode).lower()},
        )

    async def test_phase1_database_migration_preserves_unowned_run(self) -> None:
        legacy_path = Path(self.database_dir.name) / "phase1.sqlite3"
        connection = sqlite3.connect(legacy_path)
        try:
            connection.execute("""CREATE TABLE processing_runs (
                run_id TEXT PRIMARY KEY, created_at TEXT NOT NULL,
                pdf_filename TEXT NOT NULL, excel_filename TEXT NOT NULL,
                pdf_bytes BLOB NOT NULL, excel_bytes BLOB NOT NULL, result_json TEXT NOT NULL
            )""")
            connection.execute("""INSERT INTO processing_runs VALUES
                ('legacy-run','2025-01-01T00:00:00Z','legacy.pdf','legacy.xlsx',X'25504446',X'504B', '{}')""")
            connection.commit()
        finally:
            connection.close()
        with patch.dict(os.environ, {"UNTANGLE_DB_PATH": str(legacy_path)}):
            run_storage.initialize_database()
            self.assertEqual(run_storage.list_runs("demo-user"), [])
            self.assertIsNone(run_storage.get_run("legacy-run", "demo-user"))
            self.assertIsNone(run_storage.get_run_files("legacy-run", "demo-user"))
        connection = sqlite3.connect(legacy_path)
        try:
            owner = connection.execute("SELECT owner_id FROM processing_runs WHERE run_id='legacy-run'").fetchone()[0]
            pdf_bytes = connection.execute("SELECT pdf_bytes FROM processing_runs WHERE run_id='legacy-run'").fetchone()[0]
        finally:
            connection.close()
        self.assertIsNone(owner)
        self.assertEqual(pdf_bytes, b"%PDF")

    async def test_uploaded_pdf_does_not_receive_demo_equipment(self) -> None:
        response = await self._process_case("02_parallel_pumps")

        self.assertEqual(response.status_code, 200, response.text)
        node_tags = {node["tag"] for node in response.json()["nodes"]}
        self.assertTrue(node_tags)
        self.assertTrue(
            DEMO_EQUIPMENT_TAGS.isdisjoint(node_tags),
            node_tags,
        )

    async def test_authentication_and_run_ownership(self) -> None:
        invalid = await self.client.post(
            "/api/auth/login",
            json={"email": "demo@untangle.local", "password": "incorrect"},
        )
        self.assertEqual(invalid.status_code, 401)
        legacy_id = "phase1-unowned-run"
        with run_storage._connection() as connection:
            connection.execute(
                """INSERT INTO processing_runs
                   (run_id,created_at,pdf_filename,excel_filename,pdf_bytes,excel_bytes,result_json,owner_id)
                   VALUES(?,?,?,?,?,?,?,NULL)""",
                (legacy_id, "2025-01-01T00:00:00Z", "old.pdf", "old.xlsx", b"private pdf", b"private xlsx", "{}"),
            )
        self.assertIsNone(run_storage.get_run(legacy_id, "any-user"))
        self.assertIsNone(run_storage.get_run_files(legacy_id, "any-user"))
        self.assertNotIn(legacy_id, [item["run_id"] for item in run_storage.list_runs("any-user")])
        response = await self._process_case("08_ambiguous_crossing")
        self.assertEqual(response.status_code, 200, response.text)
        run_id = response.json()["run_id"]
        user_b = run_storage.create_user(
            "second@untangle.local", "second-demo-pass", full_name="Second User", role="Process Engineer"
        )
        from auth import issue_token
        token_b, _ = issue_token(user_b)
        client_b = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://test",
            headers={"Authorization": f"Bearer {token_b}"},
        )
        try:
            self.assertEqual((await client_b.get("/api/runs")).json()["runs"], [])
            self.assertEqual((await client_b.get(f"/api/runs/{run_id}")).status_code, 404)
            self.assertEqual((await client_b.get(f"/api/runs/{run_id}/files/pdf")).status_code, 404)
            self.assertEqual((await client_b.get(f"/api/runs/{run_id}/files/excel")).status_code, 404)
            clarification = response.json()["clarification_queue"][0]
            choice = {"run_id": run_id, "item_id": clarification["id"], "selected_option": clarification["options"][0]}
            self.assertEqual((await client_b.post("/api/resolve-ambiguity", json=choice)).status_code, 404)
            owner_list = await self.client.get("/api/runs")
            self.assertIn(run_id, [item["run_id"] for item in owner_list.json()["runs"]])
            self.assertEqual((await self.client.get(f"/api/runs/{run_id}")).status_code, 200)
            owner_file = await self.client.get(f"/api/runs/{run_id}/files/pdf")
            self.assertEqual(owner_file.status_code, 200)
            self.assertTrue(owner_file.content.startswith(b"%PDF"))
            self.assertEqual((await self.client.get(f"/api/runs/{run_id}/files/excel")).status_code, 200)
            self.assertEqual((await self.client.post("/api/resolve-ambiguity", json=choice)).status_code, 200)
        finally:
            await client_b.aclose()

    async def test_private_run_endpoints_reject_unauthenticated_access(self) -> None:
        response = await self._process_case("08_ambiguous_crossing")
        self.assertEqual(response.status_code, 200, response.text)
        run_id = response.json()["run_id"]
        clarification = response.json()["clarification_queue"][0]
        choice = {"run_id": run_id, "item_id": clarification["id"], "selected_option": clarification["options"][0]}
        self.client.headers.pop("Authorization")
        self.assertEqual((await self.client.get("/api/runs")).status_code, 401)
        self.assertEqual((await self.client.get(f"/api/runs/{run_id}")).status_code, 401)
        self.assertEqual((await self.client.get(f"/api/runs/{run_id}/files/pdf")).status_code, 401)
        self.assertEqual((await self.client.post("/api/resolve-ambiguity", json=choice)).status_code, 401)
        pdf_path, excel_path = _fixture_files("01_easy_linear")
        response = await self.client.post("/api/process-pid", files={
            "file_pdf": (pdf_path.name, pdf_path.read_bytes(), "application/pdf"),
            "file_excel": (excel_path.name, excel_path.read_bytes(), EXCEL_CONTENT_TYPE),
        })
        self.assertEqual(response.status_code, 401)

    async def test_authenticated_run_file_cors_for_frontend_origins(self) -> None:
        for origin in ("http://localhost:8080", "http://127.0.0.1:8080"):
            with self.subTest(origin=origin):
                preflight = await self.client.options(
                    "/api/runs/example-run/files/pdf",
                    headers={
                        "Origin": origin,
                        "Access-Control-Request-Method": "GET",
                        "Access-Control-Request-Headers": "authorization",
                    },
                )
                self.assertEqual(preflight.status_code, 200, preflight.text)
                self.assertEqual(preflight.headers.get("access-control-allow-origin"), origin)
                self.assertIn("authorization", preflight.headers.get("access-control-allow-headers", "").lower())
                self.assertNotIn("access-control-allow-credentials", preflight.headers)

                response = await self.client.get(
                    "/api/runs/example-run/files/pdf", headers={"Origin": origin}
                )
                self.assertEqual(response.status_code, 404)
                self.assertEqual(response.headers.get("access-control-allow-origin"), origin)

    async def test_scanned_pdf_reports_ocr_unavailability_without_fabricated_matches(
        self,
    ) -> None:
        response = await self._process_case("06_scanned_image_only")

        self.assertEqual(response.status_code, 400, response.text)
        detail = response.json()["detail"].lower()
        self.assertIn("no extractable text", detail)
        self.assertIn("ocr", detail)

    async def test_case_07_keeps_deliberately_absent_instrument_tags_missing(
        self,
    ) -> None:
        response = await self._process_case("07_missing_ambiguous_tags")

        self.assertEqual(response.status_code, 200, response.text)
        rows = {row["tag_name"]: row for row in response.json()["tag_audit"]}
        self.assertEqual(rows["FT-601"]["status"], "Matched")
        self.assertEqual(rows["TT-601"]["status"], "Matched")
        for tag in ("LT-609", "FT-601A", "PT-699"):
            self.assertEqual(rows[tag]["status"], "Missing in P&ID")
            self.assertEqual(rows[tag]["source_bbox"], [0, 0, 0, 0])

    async def test_excel_only_tag_is_not_reported_as_detected(self) -> None:
        pdf_path, _ = _fixture_files("01_easy_linear")
        response = await self.client.post(
            "/api/process-pid",
            files={
                "file_pdf": (
                    pdf_path.name,
                    pdf_path.read_bytes(),
                    "application/pdf",
                ),
                "file_excel": (
                    "tags.xlsx",
                    _workbook_bytes(
                        [("PT-101", "Pressure"), ("PT-999", "Pressure")]
                    ),
                    EXCEL_CONTENT_TYPE,
                ),
            },
        )

        self.assertEqual(response.status_code, 200, response.text)
        rows = {row["tag_name"]: row for row in response.json()["tag_audit"]}
        self.assertEqual(rows["PT-101"]["status"], "Matched")
        self.assertGreater(
            rows["PT-101"]["source_bbox"][2],
            rows["PT-101"]["source_bbox"][0],
        )
        self.assertEqual(rows["PT-999"]["status"], "Missing in P&ID")
        self.assertEqual(rows["PT-999"]["source_bbox"], [0, 0, 0, 0])

    async def test_invalid_excel_does_not_substitute_demo_inventory(self) -> None:
        pdf_path, _ = _fixture_files("01_easy_linear")
        response = await self.client.post(
            "/api/process-pid",
            files={
                "file_pdf": (
                    pdf_path.name,
                    pdf_path.read_bytes(),
                    "application/pdf",
                ),
                "file_excel": (
                    "invalid.xlsx",
                    b"not an Excel workbook",
                    EXCEL_CONTENT_TYPE,
                ),
            },
        )

        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("Unable to parse Excel inventory", response.json()["detail"])

    async def test_ambiguity_endpoint_does_not_create_fallback_demo_edges(self) -> None:
        response = await self.client.post(
            "/api/resolve-ambiguity",
            json={"run_id": "missing-run", "item_id": "amb_04", "selected_option": "junction"},
        )
        self.assertEqual(response.status_code, 404, response.text)

    async def test_unresolved_clarification_remains_available_for_later_resolution(self) -> None:
        response = await self._process_case("08_ambiguous_crossing")
        self.assertEqual(response.status_code, 200, response.text)
        run_id = response.json()["run_id"]
        item_id = response.json()["clarification_queue"][0]["id"]
        original_edge_ids = {edge["id"] for edge in response.json()["edges"]}

        unresolved = await self.client.post(
            "/api/resolve-ambiguity",
            json={"run_id": run_id, "item_id": item_id, "selected_option": "Unresolved / Keep Flagged"},
        )
        self.assertEqual(unresolved.status_code, 200, unresolved.text)
        self.assertFalse(unresolved.json()["topology_changed"])
        self.assertIn("remains unresolved", unresolved.json()["message"])
        self.assertEqual(unresolved.json()["kpis"]["clarification_items"], 1)
        saved = await self.client.get(f"/api/runs/{run_id}")
        self.assertEqual(saved.json()["clarification_queue"][0]["status"], "unresolved")
        self.assertEqual(
            {edge["id"] for edge in unresolved.json()["edges"]}, original_edge_ids
        )
        self.assertIn("Partially verified", unresolved.json()["kpis"]["graph_integrity"])
        self.assertNotIn("100%", unresolved.json()["kpis"]["graph_integrity"])

        resolved = await self.client.post(
            "/api/resolve-ambiguity",
            json={
                "run_id": run_id,
                "item_id": item_id,
                "selected_option": "Connected Junction (4-Way Tee)",
            },
        )
        self.assertEqual(resolved.status_code, 200, resolved.text)
        self.assertTrue(resolved.json()["topology_changed"])
        self.assertEqual(resolved.json()["kpis"]["clarification_items"], 0)
        saved = await self.client.get(f"/api/runs/{run_id}")
        self.assertEqual(
            saved.json()["clarification_decisions"][0]["selected_option"],
            "Connected Junction (4-Way Tee)",
        )
        self.assertNotEqual(
            {edge["id"] for edge in resolved.json()["edges"]}, original_edge_ids
        )

        duplicate = await self.client.post(
            "/api/resolve-ambiguity",
            json={
                "run_id": run_id,
                "item_id": item_id,
                "selected_option": "Connected Junction (4-Way Tee)",
            },
        )
        self.assertEqual(duplicate.status_code, 404, duplicate.text)

    async def test_invalid_clarification_choice_does_not_consume_item(self) -> None:
        response = await self._process_case("08_ambiguous_crossing")
        self.assertEqual(response.status_code, 200, response.text)
        run_id = response.json()["run_id"]
        item_id = response.json()["clarification_queue"][0]["id"]
        invalid = await self.client.post(
            "/api/resolve-ambiguity",
            json={"run_id": run_id, "item_id": item_id, "selected_option": "Invented choice"},
        )
        self.assertEqual(invalid.status_code, 422, invalid.text)
        valid = await self.client.post(
            "/api/resolve-ambiguity",
            json={"run_id": run_id, "item_id": item_id, "selected_option": "Connected Junction (4-Way Tee)"},
        )
        self.assertEqual(valid.status_code, 200, valid.text)

    async def test_crossing_choice_changes_graph_and_provides_source_crop(self) -> None:
        crossover_run = await self._process_case("08_ambiguous_crossing")
        junction_run = await self._process_case("08_ambiguous_crossing")
        self.assertEqual(crossover_run.status_code, 200, crossover_run.text)
        self.assertEqual(junction_run.status_code, 200, junction_run.text)
        crossover_payload = crossover_run.json()
        junction_payload = junction_run.json()
        item = crossover_payload["clarification_queue"][0]
        self.assertEqual(item["id"], junction_payload["clarification_queue"][0]["id"])
        self.assertEqual(item["source_page"], 0)
        self.assertTrue(item["evidence_crop_data_uri"].startswith("data:image/png;base64,"))
        self.assertIn("connectivity is provisional", crossover_payload["kpis"]["graph_integrity"])

        crossover = await self.client.post(
            "/api/resolve-ambiguity",
            json={"run_id": crossover_payload["run_id"], "item_id": item["id"], "selected_option": "Crossover / No Connection"},
        )
        junction = await self.client.post(
            "/api/resolve-ambiguity",
            json={"run_id": junction_payload["run_id"], "item_id": item["id"], "selected_option": "Connected Junction (4-Way Tee)"},
        )
        self.assertEqual(crossover.status_code, 200, crossover.text)
        self.assertEqual(junction.status_code, 200, junction.text)
        self.assertFalse(crossover.json()["topology_changed"])
        self.assertTrue(junction.json()["topology_changed"])

        crossover_graph = nx.Graph()
        junction_graph = nx.Graph()
        crossover_graph.add_nodes_from(node["tag"] for node in crossover.json()["nodes"])
        junction_graph.add_nodes_from(node["tag"] for node in junction.json()["nodes"])
        for edge in crossover.json()["edges"]:
            crossover_graph.add_edge(edge["source"], edge["target"])
        for edge in junction.json()["edges"]:
            junction_graph.add_edge(edge["source"], edge["target"])
        self.assertFalse(nx.has_path(crossover_graph, "P-801", "E-801"))
        self.assertTrue(nx.has_path(junction_graph, "P-801", "E-801"))
        self.assertTrue(any(node["type"] == "junction" for node in junction.json()["nodes"]))
        self.assertTrue(all(edge["flow_direction"] == "undirected" for edge in junction.json()["edges"]))
        self.assertIn("user-confirmed", junction.json()["kpis"]["graph_integrity"])
        junction_node_tags = {
            node["tag"]
            for node in junction.json()["nodes"]
            if node["type"] == "junction"
        }
        crossing_edges = [
            edge
            for edge in junction.json()["edges"]
            if edge["source"] in junction_node_tags
            or edge["target"] in junction_node_tags
        ]
        self.assertTrue(crossing_edges)
        self.assertTrue(
            all(edge["evidence_status"] == "user_confirmed_crossing" for edge in crossing_edges)
        )
        self.assertTrue(all(edge["source_segment_ids"] for edge in junction.json()["edges"]))
        saved = await self.client.get(f"/api/runs/{junction_payload['run_id']}")
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(saved.json()["nodes"], junction.json()["nodes"])
        self.assertEqual(saved.json()["edges"], junction.json()["edges"])
        self.assertEqual(
            saved.json()["clarification_decisions"][0]["selected_option"],
            "Connected Junction (4-Way Tee)",
        )

    async def test_live_upload_clarifications_are_source_evidence_only(self) -> None:
        response = await self._process_case("01_easy_linear")
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertGreater(payload["kpis"]["clarification_items"], 0)
        self.assertEqual(
            payload["kpis"]["clarification_items"],
            len(payload["clarification_queue"]),
        )
        self.assertTrue(
            all(item["id"].startswith("nozzle_") for item in payload["clarification_queue"])
        )
        self.assertTrue(
            all(
                item["evidence_crop_data_uri"].startswith("data:image/png;base64,")
                and item["options"][-1] == "Unresolved / Keep Flagged"
                for item in payload["clarification_queue"]
            )
        )

    async def test_demo_mode_does_not_add_clarifications_not_in_source(self) -> None:
        live = await self._process_case("01_easy_linear")
        response = await self._process_case("01_easy_linear", demo_mode=True)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(live.status_code, 200, live.text)
        self.assertEqual(
            response.json()["kpis"]["clarification_items"],
            len(response.json()["clarification_queue"]),
        )
        self.assertEqual(
            [item["id"] for item in response.json()["clarification_queue"]],
            [item["id"] for item in live.json()["clarification_queue"]],
        )

    async def test_supplied_nozzle_decision_changes_graph_and_persists(self) -> None:
        pdf_path, excel_path = _supplied_fixture_files("08_multistage_recycle")

        async def upload() -> dict:
            response = await self.client.post(
                "/api/process-pid",
                files={
                    "file_pdf": (pdf_path.name, pdf_path.read_bytes(), "application/pdf"),
                    "file_excel": (
                        excel_path.name,
                        excel_path.read_bytes(),
                        EXCEL_CONTENT_TYPE,
                    ),
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            return response.json()

        connected_run = await upload()
        disconnected_run = await upload()
        def p801_tag_connection(payload: dict) -> dict | None:
            return next(
                (
                    edge
                    for edge in payload["edges"]
                    if {edge["source"], edge["target"]} == {"P-801", "TK-801"}
                ),
                None,
            )

        connected_edge = p801_tag_connection(connected_run)
        self.assertIsNotNone(connected_edge)
        self.assertEqual(connected_edge["evidence_status"], "provisional_nozzle_alignment")
        item_id = next(
            item["id"]
            for item in connected_run["clarification_queue"]
            if "_p_801" in item["id"] and "d2_i0" in item["id"]
        )
        item = next(item for item in connected_run["clarification_queue"] if item["id"] == item_id)
        self.assertTrue(item["evidence_crop_data_uri"].startswith("data:image/png;base64,"))
        connected_option = "Connected to P-801"
        disconnected_option = "Not connected to P-801"
        self.assertEqual(item["options"][:2], [connected_option, disconnected_option])

        invalid = await self.client.post(
            "/api/resolve-ambiguity",
            json={"run_id": connected_run["run_id"], "item_id": item_id, "selected_option": "Connect to guessed pump"},
        )
        self.assertEqual(invalid.status_code, 422, invalid.text)
        valid = await self.client.post(
            "/api/resolve-ambiguity",
            json={"run_id": connected_run["run_id"], "item_id": item_id, "selected_option": connected_option},
        )
        self.assertEqual(valid.status_code, 200, valid.text)
        confirmed_edge = p801_tag_connection(valid.json())
        self.assertIsNotNone(confirmed_edge)
        self.assertEqual(confirmed_edge["evidence_status"], "user_confirmed_nozzle")
        saved = await self.client.get(f"/api/runs/{connected_run['run_id']}")
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(
            saved.json()["clarification_decisions"][0]["selected_option"],
            connected_option,
        )

        disconnect_id = next(
            item["id"]
            for item in disconnected_run["clarification_queue"]
            if "_p_801" in item["id"] and "d2_i0" in item["id"]
        )
        disconnected = await self.client.post(
            "/api/resolve-ambiguity",
            json={
                "run_id": disconnected_run["run_id"],
                "item_id": disconnect_id,
                "selected_option": disconnected_option,
            },
        )
        self.assertEqual(disconnected.status_code, 200, disconnected.text)
        self.assertIsNone(p801_tag_connection(disconnected.json()))
        self.assertTrue(disconnected.json()["topology_changed"])

    async def test_successful_run_persists_result_and_original_uploads(self) -> None:
        response = await self._process_case("02_parallel_pumps")
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertTrue(payload["run_id"])

        retrieved = await self.client.get(f"/api/runs/{payload['run_id']}")
        self.assertEqual(retrieved.status_code, 200, retrieved.text)
        self.assertEqual(retrieved.json(), payload)

        pdf_path, excel_path = _fixture_files("02_parallel_pumps")
        saved_pdf = await self.client.get(
            f"/api/runs/{payload['run_id']}/files/pdf"
        )
        saved_excel = await self.client.get(
            f"/api/runs/{payload['run_id']}/files/excel"
        )
        self.assertEqual(saved_pdf.content, pdf_path.read_bytes())
        self.assertEqual(saved_excel.content, excel_path.read_bytes())

        run_list = await self.client.get("/api/runs")
        self.assertEqual(run_list.status_code, 200, run_list.text)
        self.assertEqual(run_list.json()["runs"][0]["run_id"], payload["run_id"])

    async def test_unknown_run_returns_404(self) -> None:
        response = await self.client.get("/api/runs/unknown-run-id")
        self.assertEqual(response.status_code, 404, response.text)

    async def test_separate_runs_keep_their_own_results(self) -> None:
        first = await self._process_case("01_easy_linear")
        second = await self._process_case("02_parallel_pumps")
        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(second.status_code, 200, second.text)
        first_result = first.json()
        second_result = second.json()
        self.assertNotEqual(first_result["run_id"], second_result["run_id"])

        retrieved_first = await self.client.get(
            f"/api/runs/{first_result['run_id']}"
        )
        retrieved_second = await self.client.get(
            f"/api/runs/{second_result['run_id']}"
        )
        self.assertEqual(retrieved_first.json(), first_result)
        self.assertEqual(retrieved_second.json(), second_result)
        self.assertNotEqual(
            {node["tag"] for node in retrieved_first.json()["nodes"]},
            {node["tag"] for node in retrieved_second.json()["nodes"]},
        )

    async def test_clarification_decision_is_scoped_to_its_run(self) -> None:
        first = await self._process_case("08_ambiguous_crossing")
        second = await self._process_case("08_ambiguous_crossing")
        first_id = first.json()["run_id"]
        second_id = second.json()["run_id"]
        item_id = first.json()["clarification_queue"][0]["id"]

        decision = await self.client.post(
            "/api/resolve-ambiguity",
            json={
                "run_id": first_id,
                "item_id": item_id,
                "selected_option": "Connected Junction (4-Way Tee)",
            },
        )
        self.assertEqual(decision.status_code, 200, decision.text)
        self.assertEqual(decision.json()["run_id"], first_id)

        first_saved = await self.client.get(f"/api/runs/{first_id}")
        second_saved = await self.client.get(f"/api/runs/{second_id}")
        self.assertEqual(
            first_saved.json()["clarification_decisions"][0]["selected_option"],
            "Connected Junction (4-Way Tee)",
        )
        self.assertEqual(
            next(
                item["status"]
                for item in second_saved.json()["clarification_queue"]
                if item["id"] == item_id
            ),
            "pending",
        )

    async def test_audit_counters_are_defined_against_rows_and_unique_tags(self) -> None:
        response = await self._process_case("01_easy_linear")
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        counts = payload["kpis"]
        status_counts = {
            status: sum(row["status"] == status for row in payload["tag_audit"])
            for status in ("Matched", "Missing in P&ID", "Ambiguous")
        }
        self.assertEqual(counts["unique_inventory_tags"], counts["total_tags"])
        self.assertEqual(counts["inventory_rows"], counts["usable_inventory_rows"])
        self.assertEqual(counts["matched_tags"], status_counts["Matched"])
        self.assertEqual(counts["missing_tags"], status_counts["Missing in P&ID"])
        inventory_ambiguous = sum(
            row["status"] == "Ambiguous" and row["excel_row"] > 0
            for row in payload["tag_audit"]
        )
        pdf_only = sum(
            row.get("reason") == "unmatched_pdf_detection"
            for row in payload["tag_audit"]
        )
        self.assertEqual(counts["ambiguous_tags"], inventory_ambiguous)
        self.assertEqual(counts["pdf_only_tags"], pdf_only)
        self.assertEqual(counts["pending_review_tags"], status_counts.get("Pending Review", 0))
        self.assertEqual(counts["audit_rows"], len(payload["tag_audit"]))
        valid_source_rows = sum(
            has_valid_source_bbox(row["source_bbox"])
            for row in payload["tag_audit"]
        )
        self.assertEqual(counts["detected_source_tags"], valid_source_rows)
        self.assertEqual(
            counts["audit_rows"],
            counts["matched_tags"] + counts["missing_tags"] + counts["ambiguous_tags"] + counts["pdf_only_tags"] + counts["pending_review_tags"],
        )
        self.assertEqual(counts["clarification_items"], len(payload["clarification_queue"]))

    async def test_duplicate_inventory_tag_is_grouped_and_not_counted_as_matched(self) -> None:
        pdf_path, _ = _fixture_files("01_easy_linear")
        response = await self.client.post(
            "/api/process-pid",
            files={
                "file_pdf": (pdf_path.name, pdf_path.read_bytes(), "application/pdf"),
                "file_excel": (
                    "duplicates.xlsx",
                    _workbook_bytes([
                        ("PT-101", "Pressure"),
                        ("", ""),
                        ("PT-101", "Pressure"),
                    ]),
                    EXCEL_CONTENT_TYPE,
                ),
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        rows = [row for row in payload["tag_audit"] if row["tag_name"] == "PT-101"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["status"], "Ambiguous")
        self.assertEqual(rows[0]["excel_rows"], [2, 4])
        self.assertEqual(payload["kpis"]["matched_tags"], 0)
        self.assertEqual(payload["kpis"]["ambiguous_tags"], 1)
        self.assertEqual(payload["kpis"]["duplicate_inventory_tags"], 1)
        self.assertEqual(payload["kpis"]["duplicate_inventory_rows"], 1)
        self.assertEqual(payload["kpis"]["inventory_rows"], 3)
        self.assertEqual(payload["kpis"]["usable_inventory_rows"], 2)

    async def test_all_fixture_pairs_return_evidence_based_results(self) -> None:
        for case_name in (
            "01_easy_linear",
            "02_parallel_pumps",
            "03_branch_merge",
            "04_control_utility",
            "05_dense_recycle",
            "06_scanned_image_only",
            "07_missing_ambiguous_tags",
            "08_ambiguous_crossing",
            "09_bypass_recycle",
            "10_dense_shared_headers",
        ):
            with self.subTest(case=case_name):
                response = await self._process_case(case_name)
                if case_name == "06_scanned_image_only":
                    self.assertEqual(response.status_code, 400, response.text)
                    self.assertIn("OCR", response.json()["detail"])
                    continue

                self.assertEqual(response.status_code, 200, response.text)
                payload = response.json()
                self.assertNotIn("100%", payload["kpis"]["graph_integrity"])
                equipment_tags = {
                    node["tag"]
                    for node in payload["nodes"]
                    if node["type"] != "junction"
                }
                _, excel_path = _fixture_files(case_name)
                workbook = load_workbook(
                    excel_path, data_only=True, read_only=True
                )
                equipment_rows = list(workbook["Major Equipment"].values)
                headers = [str(value) for value in equipment_rows[0]]
                tag_column = headers.index("Equipment Tag")
                expected_column = headers.index("Expected in PDF")
                expected_tags = {
                    str(row[tag_column]).strip()
                    for row in equipment_rows[1:]
                    if str(row[expected_column]).strip().lower() == "yes"
                }
                self.assertEqual(equipment_tags, expected_tags)
                tag_rows = list(workbook["Instrument Tag List"].values)
                tag_headers = [str(value) for value in tag_rows[0]]
                tag_name_column = tag_headers.index("Tag Number")
                tag_expected_column = tag_headers.index("Expected in PDF")
                expected_visible_tags = {
                    str(row[tag_name_column]).strip()
                    for row in tag_rows[1:]
                    if str(row[tag_expected_column]).strip().lower() == "yes"
                }
                expected_visible_tags.update(expected_tags)
                actual_matched_tags = {
                    row["tag_name"]
                    for row in payload["tag_audit"]
                    if row["status"] == "Matched"
                }
                self.assertTrue(expected_visible_tags.issubset(actual_matched_tags))
                if case_name in {
                    "08_ambiguous_crossing",
                    "09_bypass_recycle",
                    "10_dense_shared_headers",
                }:
                    self.assertEqual(expected_visible_tags, actual_matched_tags)
                    expected_missing_tags = {
                        str(row[tag_name_column]).strip()
                        for row in tag_rows[1:]
                        if str(row[tag_expected_column]).strip().lower() == "no"
                    }
                    expected_missing_tags.update(
                        str(row[tag_column]).strip()
                        for row in equipment_rows[1:]
                        if str(row[expected_column]).strip().lower() == "no"
                    )
                    actual_missing_tags = {
                        row["tag_name"]
                        for row in payload["tag_audit"]
                        if row["status"] == "Missing in P&ID"
                    }
                    self.assertEqual(expected_missing_tags, actual_missing_tags)
                for node in payload["nodes"]:
                    self.assertEqual(node["source_page"], 0)
                    self.assertEqual(len(node["source_bbox"]), 4)
                    if node["type"] != "junction":
                        x1, y1, x2, y2 = node["source_bbox"]
                        self.assertGreater(x2, x1)
                        self.assertGreater(y2, y1)
                self.assertTrue(
                    all(edge["relation"] == "process_pipe" for edge in payload["edges"])
                )
                workbook.close()

        response = await self._process_case("08_ambiguous_crossing")
        self.assertEqual(response.status_code, 200, response.text)
        run_id = response.json()["run_id"]
        crossing_id = response.json()["clarification_queue"][0]["id"]
        response = await self.client.post(
            "/api/resolve-ambiguity",
            json={"run_id": run_id, "item_id": crossing_id, "selected_option": "Unresolved / Keep Flagged"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertNotIn("100%", response.json()["kpis"]["graph_integrity"])
        self.assertIn("flow direction is unresolved", response.json()["kpis"]["graph_integrity"])

    async def test_all_supplied_cases_use_live_upload_processing(self) -> None:
        cases = (
            "01_easy_linear",
            "02_parallel_pumps",
            "03_branch_merge",
            "04_control_utility",
            "05_dense_recycle",
            "06_scanned_image_only",
            "07_missing_ambiguous_tags",
            "08_multistage_recycle",
            "09_parallel_filter_trains",
            "10_control_loop_utility",
            "11_junction_crossover",
            "12_reactor_recirculation",
            "13_dense_multi_branch",
        )
        for case_name in cases:
            if case_name in {
                "01_easy_linear", "02_parallel_pumps", "03_branch_merge",
                "04_control_utility", "05_dense_recycle", "06_scanned_image_only",
                "07_missing_ambiguous_tags",
            }:
                pdf_path, excel_path = _fixture_files(case_name)
            else:
                pdf_path, excel_path = _supplied_fixture_files(case_name)
            with self.subTest(case=case_name):
                if not pdf_path.is_file() or not excel_path.is_file():
                    self.skipTest("Supplied mid-complex fixture pack is not present")
                response = await self.client.post(
                    "/api/process-pid",
                    files={
                        "file_pdf": (pdf_path.name, pdf_path.read_bytes(), "application/pdf"),
                        "file_excel": (
                            excel_path.name,
                            excel_path.read_bytes(),
                            EXCEL_CONTENT_TYPE,
                        ),
                    },
                )
                if case_name == "06_scanned_image_only":
                    self.assertEqual(response.status_code, 400, response.text)
                    detail = response.json()["detail"].lower()
                    self.assertIn("no extractable text", detail)
                    self.assertIn("ocr", detail)
                    continue
                self.assertEqual(response.status_code, 200, response.text)
                payload = response.json()
                audit = {row["tag_name"]: row for row in payload["tag_audit"]}
                for row in payload["tag_audit"]:
                    if row["status"] == "Matched":
                        self.assertGreater(row["source_bbox"][2], row["source_bbox"][0])
                        self.assertGreater(row["source_bbox"][3], row["source_bbox"][1])
                    if row["status"] == "Missing in P&ID":
                        self.assertEqual(row["source_bbox"], [0, 0, 0, 0])
                self.assertEqual(
                    payload["kpis"]["matched_tags"],
                    sum(row["status"] == "Matched" for row in payload["tag_audit"]),
                )
                self.assertEqual(
                    payload["kpis"]["missing_tags"],
                    sum(row["status"] == "Missing in P&ID" for row in payload["tag_audit"]),
                )
                self.assertEqual(
                    payload["kpis"]["ambiguous_tags"],
                    sum(
                        row["status"] == "Ambiguous" and row["excel_row"] > 0
                        for row in payload["tag_audit"]
                    ),
                )
                self.assertEqual(payload["kpis"]["audit_rows"], len(payload["tag_audit"]))
                self.assertTrue(
                    all(edge["relation"] == "process_pipe" for edge in payload["edges"])
                )
                self.assertTrue(payload["edges"], f"{case_name}: source linework yielded no graph edges")
                self.assertTrue(
                    all(edge["source_segment_ids"] for edge in payload["edges"]),
                    f"{case_name}: graph edge lacks source segment evidence",
                )
                self.assertFalse("100%" in payload["kpis"]["graph_integrity"])

                deliberately_absent = {
                    "07_missing_ambiguous_tags": {"LT-609", "FT-601A", "PT-699", "E-609"},
                    "08_multistage_recycle": {"PT-899"},
                    "10_control_loop_utility": {"LSH-1099"},
                    "12_reactor_recirculation": {"PSV-1201"},
                    "13_dense_multi_branch": {"FT-1301A", "XV-1399"},
                }.get(case_name, set())
                for tag in deliberately_absent:
                    self.assertEqual(audit[tag]["status"], "Missing in P&ID")
                    self.assertEqual(audit[tag]["source_bbox"], [0, 0, 0, 0])
                if case_name == "09_parallel_filter_trains":
                    self.assertEqual(payload["kpis"]["matched_tags"], payload["kpis"]["total_tags"])


class SourceTopologyTests(unittest.TestCase):
    def _fixture_result(self, case_name: str) -> tuple[dict, dict]:
        pdf_path, excel_path = _fixture_files(case_name)
        pdf_result = process_pdf(pdf_path.read_bytes(), filename=pdf_path.name)
        workbook = load_workbook(excel_path, data_only=True, read_only=True)
        rows = list(workbook["Instrument Tag List"].values)
        headers = [str(value) for value in rows[0]]
        tag_column = headers.index("Tag Number")
        type_column = headers.index("Instrument Type")
        inventory = [
            {
                "tag_name": str(row[tag_column]).strip(),
                "type": str(row[type_column] or "Instrument"),
                "excel_row": index + 1,
            }
            for index, row in enumerate(rows[1:], start=1)
            if row[tag_column]
        ]
        audit = match_excel_inventory(pdf_result["tags"], inventory)
        result = build_process_graph(
            pdf_result["equipment"], audit, pdf_result["drawing_evidence"]
        )
        workbook.close()
        return pdf_result, result

    def test_right_side_legend_and_aligned_nozzles_restore_fixture_topology(self) -> None:
        pdf_path, _ = _supplied_fixture_files("08_multistage_recycle")
        pdf_result = process_pdf(pdf_path.read_bytes(), filename=pdf_path.name)
        detected_equipment = {item["tag"]: item for item in pdf_result["equipment"]}
        expected_visible = {
            "TK-801", "P-801", "F-801", "E-801", "V-801", "P-802", "FV-801", "M-801"
        }
        self.assertTrue(expected_visible.issubset(detected_equipment))
        self.assertEqual(
            detected_equipment["F-801"]["type"],
            "vessel",
            "the visible source descriptor classifies F-801 as a vessel",
        )
        self.assertGreater(len(pdf_result["drawing_evidence"]["segments"]), 10)
        page_height = pdf_result["drawing_evidence"]["page_bounds"]["0"][3]
        self.assertTrue(
            all(segment["source_bbox"][1] < page_height * 0.9 for segment in pdf_result["drawing_evidence"]["segments"]),
            "the process legend swatch must not be treated as process piping",
        )

        graph_result = build_process_graph(
            pdf_result["equipment"], [], pdf_result["drawing_evidence"]
        )
        graph = graph_result["nx_graph"]
        self.assertTrue(graph_result["edges"])
        self.assertTrue(nx.has_path(graph, "TK-801", "V-801"))
        self.assertTrue(nx.has_path(graph, "P-801", "V-801"))
        self.assertTrue(nx.cycle_basis(graph), "the source-supported bypass loop must remain visible")
        self.assertTrue(
            all(edge["source_segment_ids"] for edge in graph_result["edges"])
        )
        self.assertTrue(
            all(edge["flow_direction"] == "undirected" for edge in graph_result["edges"])
        )
        self.assertTrue(
            any(edge["evidence_status"] == "provisional_nozzle_alignment" for edge in graph_result["edges"])
        )
        nozzle_items = [
            item for item in pdf_result["clarifications"] if item["id"].startswith("nozzle_")
        ]
        self.assertTrue(nozzle_items)
        self.assertTrue(
            all(item["evidence_crop_data_uri"].startswith("data:image/png;base64,") for item in nozzle_items)
        )
        self.assertTrue(
            all(item["options"][-1] == "Unresolved / Keep Flagged" for item in nozzle_items)
        )

    def test_all_native_fixture_connections_are_supported_by_source_paths(self) -> None:
        for case_name in (
            "01_easy_linear",
            "02_parallel_pumps",
            "03_branch_merge",
            "04_control_utility",
            "05_dense_recycle",
            "07_missing_ambiguous_tags",
            "08_ambiguous_crossing",
            "09_bypass_recycle",
            "10_dense_shared_headers",
        ):
            with self.subTest(case=case_name):
                pdf_path, excel_path = _fixture_files(case_name)
                pdf_result = process_pdf(pdf_path.read_bytes(), filename=pdf_path.name)
                workbook = load_workbook(excel_path, data_only=True, read_only=True)
                tag_rows = list(workbook["Instrument Tag List"].values)
                tag_headers = [str(value) for value in tag_rows[0]]
                tag_column = tag_headers.index("Tag Number")
                type_column = tag_headers.index("Instrument Type")
                inventory = [
                    {
                        "tag_name": str(row[tag_column]).strip(),
                        "type": str(row[type_column] or "Instrument"),
                        "excel_row": index + 1,
                    }
                    for index, row in enumerate(tag_rows[1:], start=1)
                    if row[tag_column]
                ]
                graph = build_process_graph(
                    pdf_result["equipment"],
                    match_excel_inventory(pdf_result["tags"], inventory),
                    pdf_result["drawing_evidence"],
                )
                workbook_connections = list(
                    workbook["Expected Connections"].values
                )
                headers = [str(value) for value in workbook_connections[0]]
                from_column = headers.index("From")
                to_column = headers.index("To")
                relationship_column = headers.index("Relationship")
                expected = [
                    row
                    for row in workbook_connections[1:]
                    if str(row[relationship_column]).lower()
                    not in {"utility", "instrument-signal"}
                ]
                for connection in expected:
                    source = str(connection[from_column]).strip()
                    target = str(connection[to_column]).strip()
                    self.assertTrue(
                        nx.has_path(graph["nx_graph"], source, target),
                        f"{case_name}: no source-supported route {source} -> {target}",
                    )
                self.assertTrue(
                    all(edge["relation"] == "process_pipe" for edge in graph["edges"])
                )
                self.assertTrue(
                    all(edge["flow_direction"] == "undirected" for edge in graph["edges"])
                )
                self.assertNotIn("100%", graph["graph_integrity"])
                workbook.close()

    def test_parallel_pumps_branches_and_merges_remain_distinct(self) -> None:
        _, graph = self._fixture_result("02_parallel_pumps")
        tags = {node["tag"] for node in graph["nodes"]}
        self.assertTrue({"P-201A", "P-201B"}.issubset(tags))
        self.assertGreaterEqual(
            sum(node["type"] == "junction" for node in graph["nodes"]), 2
        )
        self.assertTrue(
            any(node["tag"] == "P-201A" for node in graph["nodes"])
        )

    def test_branch_merge_and_recycle_loops_are_retained(self) -> None:
        _, branch_graph = self._fixture_result("03_branch_merge")
        self.assertGreaterEqual(
            sum(node["type"] == "junction" for node in branch_graph["nodes"]), 2
        )
        _, recycle_graph = self._fixture_result("05_dense_recycle")
        self.assertTrue(nx.cycle_basis(recycle_graph["nx_graph"]))
        recycle_junctions = {
            node["tag"]
            for node in recycle_graph["nodes"]
            if node["type"] == "junction"
            and recycle_graph["nx_graph"].has_edge("V-501", node["tag"])
            and recycle_graph["nx_graph"].has_edge(node["tag"], "FCV-501")
        }
        self.assertTrue(recycle_junctions)

    def test_bypass_recycle_and_dense_shared_headers_are_retained(self) -> None:
        _, bypass_graph = self._fixture_result("09_bypass_recycle")
        self.assertTrue(nx.cycle_basis(bypass_graph["nx_graph"]))
        _, dense_graph = self._fixture_result("10_dense_shared_headers")
        self.assertTrue(
            {"P-1001A", "P-1001B", "E-1001", "F-1001", "V-1001"}.issubset(
                {node["tag"] for node in dense_graph["nodes"]}
            )
        )
        self.assertTrue(all(edge["flow_direction"] == "undirected" for edge in dense_graph["edges"]))
        source_segments = [
            segment
            for segment in self._fixture_result("10_dense_shared_headers")[0]["drawing_evidence"]["segments"]
            if segment["kind"] == "process"
        ]
        self.assertGreater(len(source_segments), len(dense_graph["edges"]))

    def test_crossing_fixture_has_no_unsupported_cross_path(self) -> None:
        pdf_result, default_graph = self._fixture_result("08_ambiguous_crossing")
        crossings = find_ambiguous_process_crossings(pdf_result["drawing_evidence"])
        self.assertEqual(len(crossings), 1)
        self.assertFalse(nx.has_path(default_graph["nx_graph"], "P-801", "E-801"))

        _, excel_path = _fixture_files("08_ambiguous_crossing")
        workbook = load_workbook(excel_path, data_only=True, read_only=True)
        rows = list(workbook["Instrument Tag List"].values)
        headers = [str(value) for value in rows[0]]
        tag_column = headers.index("Tag Number")
        type_column = headers.index("Instrument Type")
        inventory = [
            {"tag_name": str(row[tag_column]).strip(), "type": str(row[type_column]), "excel_row": index + 1}
            for index, row in enumerate(rows[1:], start=1)
            if row[tag_column]
        ]
        junction_graph = build_process_graph(
            pdf_result["equipment"],
            match_excel_inventory(pdf_result["tags"], inventory),
            pdf_result["drawing_evidence"],
            crossing_resolutions={crossings[0]["id"]: CONNECTED_JUNCTION},
        )
        self.assertTrue(nx.has_path(junction_graph["nx_graph"], "P-801", "E-801"))
        workbook.close()

    def test_utility_and_instrument_signal_strokes_are_not_process_edges(self) -> None:
        pdf_result, graph = self._fixture_result("04_control_utility")
        kinds = {segment["kind"] for segment in pdf_result["drawing_evidence"]["segments"]}
        self.assertTrue({"process", "utility", "instrument_signal"}.issubset(kinds))
        self.assertTrue(all(edge["relation"] == "process_pipe" for edge in graph["edges"]))

    def test_ordering_without_pipe_geometry_cannot_create_edges(self) -> None:
        equipment = [
            {
                "tag": "TK-1",
                "type": "tank",
                "source_bbox": [0, 0, 10, 10],
            },
            {
                "tag": "V-1",
                "type": "vessel",
                "source_bbox": [100, 0, 110, 10],
            },
        ]
        result = build_process_graph(equipment, [], None)
        self.assertEqual(result["edges"], [])
        self.assertIn("Unverified", result["graph_integrity"])

    def test_near_line_endpoints_join_and_keep_source_segment_evidence(self) -> None:
        equipment = [
            {
                "tag": "P-1",
                "type": "pump",
                "source_page": 0,
                "source_bbox": [8, 48, 12, 52],
                "source_symbol_bbox": [5, 45, 15, 55],
            },
            {
                "tag": "E-1",
                "type": "exchanger",
                "source_page": 0,
                "source_bbox": [88, 48, 92, 52],
                "source_symbol_bbox": [85, 45, 95, 55],
            },
        ]
        evidence = {
            "classified_process_lines": 2,
            "page_bounds": {"0": [0, 0, 1000, 1000]},
            "segments": [
                {"id": "suction", "page": 0, "kind": "process", "start": [15, 50], "end": [50, 50]},
                {"id": "discharge", "page": 0, "kind": "process", "start": [51.25, 50], "end": [85, 50]},
            ],
        }

        result = build_process_graph(equipment, [], evidence)

        self.assertTrue(nx.has_path(result["nx_graph"], "P-1", "E-1"))
        self.assertEqual(set(result["edges"][0]["source_segment_ids"]), {"suction", "discharge"})
        self.assertEqual(result["edges"][0]["evidence_status"], "source_geometry")
        self.assertNotIn("confidence", result["edges"][0])

    def test_endpoint_gap_beyond_page_tolerance_does_not_create_a_connection(self) -> None:
        equipment = [
            {"tag": "P-1", "type": "pump", "source_page": 0, "source_bbox": [8, 48, 12, 52], "source_symbol_bbox": [5, 45, 15, 55]},
            {"tag": "E-1", "type": "exchanger", "source_page": 0, "source_bbox": [88, 48, 92, 52], "source_symbol_bbox": [85, 45, 95, 55]},
        ]
        evidence = {
            "classified_process_lines": 2,
            "page_bounds": {"0": [0, 0, 1000, 1000]},
            "segments": [
                {"id": "first", "page": 0, "kind": "process", "start": [15, 50], "end": [50, 50]},
                {"id": "second", "page": 0, "kind": "process", "start": [55, 50], "end": [85, 50]},
            ],
        }

        result = build_process_graph(equipment, [], evidence)

        self.assertFalse(nx.has_path(result["nx_graph"], "P-1", "E-1"))
        self.assertEqual(result["edges"], [])

    def test_endpoint_tolerance_scales_with_pdf_page_dimensions(self) -> None:
        from graph_engine import _endpoint_tolerance

        self.assertLess(
            _endpoint_tolerance([0, 0, 500, 500]),
            _endpoint_tolerance([0, 0, 2000, 2000]),
        )
        self.assertLessEqual(_endpoint_tolerance([0, 0, 20000, 20000]), 2.0)

    def test_duplicate_equipment_and_self_edges_are_rejected(self) -> None:
        pdf_path, _ = _fixture_files("01_easy_linear")
        pdf_result = process_pdf(pdf_path.read_bytes(), filename=pdf_path.name)
        equipment = pdf_result["equipment"] + [dict(pdf_result["equipment"][0])]
        result = build_process_graph(
            equipment, [], pdf_result["drawing_evidence"]
        )
        tags = [node["tag"] for node in result["nodes"]]
        self.assertEqual(len(tags), len(set(tags)))
        self.assertTrue(all(edge["source"] != edge["target"] for edge in result["edges"]))

    def test_unmarked_crossing_does_not_join_process_paths(self) -> None:
        equipment = [
            {"tag": "P-1", "type": "pump", "source_page": 0, "source_bbox": [8, 38, 12, 42], "source_symbol_bbox": [5, 35, 15, 45]},
            {"tag": "V-1", "type": "vessel", "source_page": 0, "source_bbox": [88, 38, 92, 42], "source_symbol_bbox": [85, 35, 95, 45]},
            {"tag": "E-1", "type": "exchanger", "source_page": 0, "source_bbox": [48, 8, 52, 12], "source_symbol_bbox": [45, 5, 55, 15]},
            {"tag": "TK-2", "type": "tank", "source_page": 0, "source_bbox": [48, 88, 52, 92], "source_symbol_bbox": [45, 85, 55, 95]},
        ]
        evidence = {
            "classified_process_lines": 2,
            "segments": [
                {"page": 0, "start": [0, 40], "end": [100, 40], "kind": "process"},
                {"page": 0, "start": [50, 0], "end": [50, 100], "kind": "process"},
            ],
        }
        result = build_process_graph(equipment, [], evidence)
        self.assertFalse(nx.has_path(result["nx_graph"], "P-1", "E-1"))
        self.assertFalse(any(node["type"] == "junction" for node in result["nodes"]))

    def test_clarification_does_not_claim_graph_validation(self) -> None:
        edge = {
            "id": "edge_a_b",
            "source": "A",
            "target": "B",
            "flow_direction": "undirected",
            "relation": "process_pipe",
            "confidence": 0.85,
        }
        result = apply_ambiguity_resolution(
            [edge],
            item_id="amb_04",
            selected_option="Connected Junction",
            retained_equipment=2,
            total_tags=1,
            matched_tags=1,
        )
        self.assertEqual(result["edges"], [edge])
        self.assertNotIn("100%", result["kpis"]["graph_integrity"])


if __name__ == "__main__":
    unittest.main()
